"""质量验证工具 - 对样例主题真实调用 LLM 生成课件，量化内容质量并做结构断言。

用法:
    python eval/quality_check.py [math|cs|general|all]

全部通过时 exit 0，否则 exit 1。
"""
from __future__ import annotations

import os
import sys
import time
from collections import Counter
from typing import Any

# ---------------------------------------------------------------------------
# 路径设置：使 core/services 可导入
# ---------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
sys.path.insert(0, BACKEND)

# ---------------------------------------------------------------------------
# 手动加载 .env（逐行解析，不依赖 python-dotenv）
# ---------------------------------------------------------------------------
_ENV_FILE = os.path.join(BACKEND, ".env")
if os.path.isfile(_ENV_FILE):
    with open(_ENV_FILE, encoding="utf-8") as _f:
        for _line in _f:
            _line = _line.strip()
            if not _line or _line.startswith("#"):
                continue
            if "=" in _line:
                _k, _v = _line.split("=", 1)
                os.environ.setdefault(_k.strip(), _v.strip())

# ---------------------------------------------------------------------------
# 特性开关
# ---------------------------------------------------------------------------
os.environ["TWO_STAGE_GEN"] = "1"
os.environ["PPTXGENJS_SERVICE_URL"] = ""   # 禁用外部服务，走 python-pptx fallback

# ---------------------------------------------------------------------------
# 输出目录
# ---------------------------------------------------------------------------
OUTPUTS_DIR = os.path.join(BACKEND, "outputs")
os.makedirs(OUTPUTS_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# 3 个样例 intent
# ---------------------------------------------------------------------------
SAMPLE_INTENTS: dict[str, dict[str, Any]] = {
    "math": {
        "topic": "高等数学 - 拉格朗日中值定理",
        "teaching_goal": "理解拉格朗日中值定理的表述、几何意义与证明，并能应用于微分估算",
        "audience": "大学本科一年级",
        "difficulty_focus": "定理证明过程与辅助函数构造",
        "key_points": [
            "中值定理表述：[a,b]连续、(a,b)可微条件下存在 c 使 f'(c)=(f(b)-f(a))/(b-a)",
            "几何意义：存在切线平行于割线的点",
            "证明思路：构造辅助函数 F(x)=f(x)-f(a)-[(f(b)-f(a))/(b-a)](x-a)，应用 Rolle 定理",
            "典型应用：函数单调性判定、不等式估算、近似计算",
        ],
        "duration": "45分钟",
        "style": "严谨学术",
    },
    "cs": {
        "topic": "数据结构 - 红黑树插入",
        "teaching_goal": "掌握红黑树五条性质以及插入后旋转与变色的修复策略",
        "audience": "大学本科二年级",
        "difficulty_focus": "插入修复的三种 case 判断与旋转变色顺序",
        "key_points": [
            "红黑树性质：节点红/黑、根黑、叶(NIL)黑、红节点子必黑、任意路径黑高相同",
            "插入基本流程：按 BST 插入后着红色，再自底向上修复违规",
            "插入修复三 case：叔红->变色上升；叔黑+zig-zag->先旋转变直线；叔黑+直线->旋转根变色",
            "旋转操作：左旋/右旋 O(1)，保持 BST 有序性",
            "时间复杂度：插入 O(log n)，修复最多 O(log n) 次变色、O(1) 次旋转",
        ],
        "duration": "50分钟",
        "style": "图示+伪代码",
    },
    "general": {
        "topic": "计算机网络 - IPv6 路由原理",
        "teaching_goal": "理解 IPv6 地址结构、NDP 邻居发现、最长前缀匹配路由及 IPv4/IPv6 过渡技术",
        "audience": "大学本科三年级",
        "difficulty_focus": "最长前缀匹配查找与过渡技术选型",
        "key_points": [
            "IPv6 地址结构：128 位，8 组 16 位十六进制，支持前导零省略与连续零段 :: 压缩",
            "NDP 邻居发现协议：替代 ARP，使用 ICMPv6 NS/NA 消息实现地址解析与重复地址检测",
            "最长前缀匹配：查路由表时选最长匹配前缀，LPM 算法可用 Trie 或 TCAM 硬件加速",
            "过渡技术：双栈、隧道（6in4/ISATAP/Teredo）、NAT64+DNS64 三类方案对比",
        ],
        "duration": "45分钟",
        "style": "简洁学术",
    },
}

# ---------------------------------------------------------------------------
# 内容指标计算
# ---------------------------------------------------------------------------

def _compute_content_metrics(slides_json: dict[str, Any]) -> dict[str, Any]:
    """从 slides_json 中计算内容质量指标。"""
    pages: list[dict] = slides_json.get("pages", []) if isinstance(slides_json, dict) else []

    page_type_counter: Counter[str] = Counter()
    bullets_total = 0
    bullet_chars_total = 0
    content_chars_total = 0

    for page in pages:
        if not isinstance(page, dict):
            continue
        ptype = str(page.get("type", "unknown"))
        page_type_counter[ptype] += 1

        # 累计各类文本
        for field in ("bullets", "takeaways", "items", "steps"):
            items = page.get(field)
            if isinstance(items, list):
                for item in items:
                    text = str(item).strip()
                    if text:
                        bullets_total += 1
                        bullet_chars_total += len(text)
                        content_chars_total += len(text)

        # code / formula 等也计入 content_chars
        for field in ("code", "explanation", "problem", "answer", "notes", "tip"):
            val = page.get(field)
            if val:
                content_chars_total += len(str(val).strip())

        for formula in (page.get("formulas") or []):
            if isinstance(formula, dict):
                content_chars_total += len(str(formula.get("expr", "")).strip())
                content_chars_total += len(str(formula.get("explanation", "")).strip())

        for side in ("left", "right"):
            col = page.get(side)
            if isinstance(col, dict):
                for item in (col.get("points") or []):
                    text = str(item).strip()
                    if text:
                        bullets_total += 1
                        bullet_chars_total += len(text)
                        content_chars_total += len(text)

    n_pages = len(pages)
    # 只统计内容页的要点均值
    content_page_count = page_type_counter.get("content", 0) + page_type_counter.get("formula", 0) + \
                         page_type_counter.get("code", 0) + page_type_counter.get("example", 0) + \
                         page_type_counter.get("two_column", 0)

    avg_bullets = round(bullets_total / content_page_count, 2) if content_page_count else 0
    avg_bullet_chars = round(bullet_chars_total / bullets_total, 2) if bullets_total else 0

    return {
        "pages": n_pages,
        "content_page_count": content_page_count,
        "bullets_total": bullets_total,
        "avg_bullets_per_page": avg_bullets,
        "avg_bullet_chars": avg_bullet_chars,
        "total_content_chars": content_chars_total,
        "page_type_dist": dict(page_type_counter),
    }


# ---------------------------------------------------------------------------
# PPTX 结构断言
# ---------------------------------------------------------------------------

def _assert_pptx_structure(path: str) -> dict[str, Any]:
    """打开 pptx，统计结构并做基本断言。返回统计字典，断言失败抛 AssertionError。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    assert os.path.isfile(path), f"pptx 文件不存在: {path}"
    size = os.path.getsize(path)
    assert size > 0, f"pptx 文件为空: {path}"

    prs = Presentation(path)
    slide_count = len(prs.slides)
    assert slide_count > 0, "pptx 中没有任何幻灯片"

    picture_count = 0
    table_count = 0
    text_shape_count = 0

    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
            except Exception:
                pass
            try:
                if shape.has_table:
                    table_count += 1
            except Exception:
                pass
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    text_shape_count += 1
            except Exception:
                pass

    return {
        "slide_count": slide_count,
        "file_size_bytes": size,
        "picture_count": picture_count,
        "table_count": table_count,
        "text_shape_count": text_shape_count,
    }


# ---------------------------------------------------------------------------
# 单主题运行
# ---------------------------------------------------------------------------

def run_topic(key: str) -> tuple[bool, dict[str, Any]]:
    """对一个样例主题完整运行：生成 slides_json -> 渲染 pptx -> 统计指标 -> 断言。

    Returns:
        (passed: bool, report: dict)
    """
    from core.two_stage_gen import generate_slides_json_two_stage
    from core.ppt_gen import generate_pptx_from_slides_json

    intent = SAMPLE_INTENTS[key]
    topic = intent["topic"]
    # 用 ASCII 安全的短 key 作文件名
    safe_name = key
    output_path = os.path.join(OUTPUTS_DIR, f"_qc_{safe_name}.pptx")

    report: dict[str, Any] = {"topic": topic, "key": key}

    try:
        # ---- Stage 1+2: 生成 slides_json ----
        t0 = time.time()
        print(f"\n[{key}] 开始生成 slides_json ...")
        slides_json = generate_slides_json_two_stage(intent, [])
        gen_elapsed = round(time.time() - t0, 2)
        print(f"[{key}] slides_json 生成完成，耗时 {gen_elapsed}s")

        # ---- 内容指标 ----
        metrics = _compute_content_metrics(slides_json)
        report["metrics"] = metrics
        report["gen_elapsed_s"] = gen_elapsed

        # ---- 渲染 pptx ----
        t1 = time.time()
        print(f"[{key}] 渲染 pptx ...")
        generate_pptx_from_slides_json(intent, slides_json, output_path)
        render_elapsed = round(time.time() - t1, 2)
        print(f"[{key}] pptx 渲染完成，耗时 {render_elapsed}s -> {output_path}")
        report["render_elapsed_s"] = render_elapsed
        report["output_path"] = output_path

        # ---- 结构断言 ----
        struct = _assert_pptx_structure(output_path)
        report["pptx_structure"] = struct

        report["passed"] = True

    except Exception as exc:
        report["passed"] = False
        report["error"] = str(exc)
        import traceback
        report["traceback"] = traceback.format_exc()

    return report["passed"], report


# ---------------------------------------------------------------------------
# 打印单主题报告
# ---------------------------------------------------------------------------

def print_report(report: dict[str, Any]) -> None:
    """打印单主题的详细报告到 stdout。"""
    key = report.get("key", "?")
    topic = report.get("topic", "?")
    passed = report.get("passed", False)
    status = "PASS" if passed else "FAIL"

    print(f"\n{'='*60}")
    print(f"[{status}] {topic}  (key={key})")
    print(f"{'='*60}")

    if not passed:
        print(f"  ERROR: {report.get('error', 'unknown')}")
        tb = report.get("traceback")
        if tb:
            print("  Traceback (last 10 lines):")
            for line in tb.strip().splitlines()[-10:]:
                print(f"    {line}")
        return

    # 生成耗时
    print(f"  gen_elapsed_s   : {report.get('gen_elapsed_s', '?')}")
    print(f"  render_elapsed_s: {report.get('render_elapsed_s', '?')}")
    print(f"  output_path     : {report.get('output_path', '?')}")

    # 内容指标
    m = report.get("metrics", {})
    print("\n  [内容指标]")
    print(f"    pages                : {m.get('pages', 0)}")
    print(f"    content_page_count   : {m.get('content_page_count', 0)}")
    print(f"    bullets_total        : {m.get('bullets_total', 0)}")
    print(f"    avg_bullets_per_page : {m.get('avg_bullets_per_page', 0)}")
    print(f"    avg_bullet_chars     : {m.get('avg_bullet_chars', 0)}")
    print(f"    total_content_chars  : {m.get('total_content_chars', 0)}")
    print(f"    page_type_dist       : {m.get('page_type_dist', {})}")

    # pptx 结构
    s = report.get("pptx_structure", {})
    print("\n  [pptx 结构]")
    print(f"    slide_count      : {s.get('slide_count', 0)}")
    print(f"    file_size_bytes  : {s.get('file_size_bytes', 0)}")
    print(f"    picture_count    : {s.get('picture_count', 0)}")
    print(f"    table_count      : {s.get('table_count', 0)}")
    print(f"    text_shape_count : {s.get('text_shape_count', 0)}")


# ---------------------------------------------------------------------------
# CLI 入口
# ---------------------------------------------------------------------------

def main() -> int:
    """CLI 入口。返回 exit code（0=全部 PASS，1=有 FAIL）。"""
    import argparse

    parser = argparse.ArgumentParser(
        description="FastPPT 质量验证工具 - 真实 LLM 调用 + pptx 结构断言"
    )
    parser.add_argument(
        "topics",
        nargs="?",
        default="all",
        choices=["math", "cs", "general", "all"],
        help="要验证的主题 (默认 all)",
    )
    args = parser.parse_args()

    if args.topics == "all":
        keys = list(SAMPLE_INTENTS.keys())
    else:
        keys = [args.topics]

    print(f"\nFastPPT 质量验证工具  topics={keys}")
    print(f"BACKEND={BACKEND}")
    print(f"OUTPUTS_DIR={OUTPUTS_DIR}")

    results: list[dict[str, Any]] = []
    for key in keys:
        passed, report = run_topic(key)
        print_report(report)
        results.append(report)

    # 总览
    print(f"\n{'='*60}")
    print("总览")
    print(f"{'='*60}")
    all_passed = True
    for r in results:
        status = "PASS" if r.get("passed") else "FAIL"
        print(f"  [{status}] {r.get('topic', r.get('key', '?'))}")
        if not r.get("passed"):
            all_passed = False

    if all_passed:
        print("\n全部 PASS")
        return 0
    else:
        print("\n存在 FAIL，请检查上方报告")
        return 1


if __name__ == "__main__":
    sys.exit(main())
