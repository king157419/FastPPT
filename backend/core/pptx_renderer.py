"""Deterministic python-pptx renderer: draws slides faithfully from slides_json.pages[].

Design goals (university-lecture quality, not generic AI slides):
* A restrained light "academic" system — one accent colour, clear type scale,
  generous margins, consistent header/footer furniture and page numbers.
* Titles are cleaned (no raw math/over-long sentences), capped to a fixed two-line
  header zone, so they never collide with the rule or the body.
* Body content is vertically centred in the content zone so pages don't look
  top-heavy/empty.
* Formulas render via matplotlib mathtext; code gets lightweight syntax
  highlighting on a dark panel; two-column / example / summary get real layouts.
Pure python-pptx (no external service). Routed behind NEW_RENDERER (default on).
"""
from __future__ import annotations

import base64
import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from uuid import uuid4

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

matplotlib.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
matplotlib.rcParams["axes.unicode_minus"] = False

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
SLIDE_W = Inches(SLIDE_W_IN)
SLIDE_H = Inches(SLIDE_H_IN)
EMU_PER_INCH = 914400

FONT_CN = "Microsoft YaHei"
FONT_MONO = "Consolas"

# --- palette (clean academic) ---
C_BG = "#FFFFFF"
C_INK = "#1C2A39"      # titles / strong text
C_BODY = "#33414F"     # body text
C_ACCENT = "#0F766E"   # brand teal — used sparingly
C_ACCENT_SOFT = "#0F766E"
C_MUTED = "#8A97A3"    # captions / footer / page number
C_HAIRLINE = "#E4E9EC"
C_SOFT = "#F1F6F5"     # subtle panel fill (teal-tint)
C_SOFT2 = "#EAF1F0"
# code panel (VS Code-ish dark)
C_CODE_BG = "#1E1E2E"
C_CODE_DEFAULT = "#D4D4D4"
C_CODE_KW = "#569CD6"
C_CODE_STR = "#CE9178"
C_CODE_NUM = "#B5CEA8"
C_CODE_COM = "#6A9955"
C_CODE_FN = "#DCDCAA"

# --- layout grid (inches) ---
MARGIN_L = 0.92
MARGIN_R = 0.92
CONTENT_W = SLIDE_W_IN - MARGIN_L - MARGIN_R
HEADER_TOP = 0.62
TITLE_TOP = 0.96
RULE_Y = 1.78          # fixed: header reserves 2 title lines, rule sits below
CONTENT_TOP = 2.02
CONTENT_BOTTOM = 6.82  # above footer
FOOTER_Y = 6.95


def render_enabled() -> bool:
    raw = os.environ.get("NEW_RENDERER", "1").strip().lower()
    return raw not in {"0", "false", "off", "no"}


def _rgb(hexstr: str | None, default: str = "#000000") -> RGBColor:
    s = (hexstr or default).lstrip("#")
    if len(s) != 6:
        s = default.lstrip("#")
    try:
        return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
    except ValueError:
        return RGBColor(0, 0, 0)


@dataclass
class RenderCtx:
    tmpdir: str
    accent: RGBColor
    ink: RGBColor
    course: str


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def _set_font(run, size=None, bold=None, italic=None, color=None, name=FONT_CN) -> None:
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if name:
        f.name = name
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", name)


def _rect(slide, left, top, width, height, color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape


def _round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    try:
        shape.shadow.inherit = False
        shape.adjustments[0] = 0.06
    except Exception:
        pass
    return shape


def _textbox(slide, left, top, width, height, *, anchor=None, wrap=True):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    return box, tf


def _set_para(p, *, align=None, space_before=None, space_after=None, line=None):
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)
    if line is not None:
        p.line_spacing = line


def _simple_text(slide, text, left, top, width, height, *, size=18, bold=False, italic=False,
                 color=None, align=PP_ALIGN.LEFT, name=FONT_CN, anchor=None):
    box, tf = _textbox(slide, left, top, width, height, anchor=anchor)
    p = tf.paragraphs[0]
    _set_para(p, align=align)
    run = p.add_run()
    run.text = text or ""
    _set_font(run, size=size, bold=bold, italic=italic, color=color or _rgb(C_BODY), name=name)
    return box


# --------------------------------------------------------------------------- #
# title cleaning
# --------------------------------------------------------------------------- #
_MATH_IN_TITLE = re.compile(r"[=^_]|f'\(|\)/\(|\\|≤|≥|∈|∑|∫|√")


def _short_title(title: str) -> str:
    """Keep titles short & math-free so the header never overflows/collides."""
    t = re.sub(r"\s+", " ", str(title or "").strip())
    if len(t) <= 22 and not _MATH_IN_TITLE.search(t):
        return t
    for sep in ("：", ":"):
        if sep in t:
            head = t.split(sep, 1)[0].strip()
            if 2 <= len(head) <= 18:
                return head
    if _MATH_IN_TITLE.search(t):
        cut = re.split(r"[=（(]|f'\(|[，,]", t)[0].strip(" ，,：:")
        if 4 <= len(cut) <= 24:
            return cut
    return t[:24].rstrip() + ("…" if len(t) > 24 else "")


def _title_pt(title: str) -> int:
    n = len(title)
    if n <= 14:
        return 30
    if n <= 22:
        return 27
    return 24


# --------------------------------------------------------------------------- #
# header / footer furniture
# --------------------------------------------------------------------------- #
def _header(slide, ctx: RenderCtx, kicker: str, title: str) -> None:
    title = _short_title(title)
    if kicker:
        box, tf = _textbox(slide, MARGIN_L, HEADER_TOP, CONTENT_W, 0.3)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = kicker
        _set_font(run, size=12.5, bold=True, color=ctx.accent)
    box, tf = _textbox(slide, MARGIN_L, TITLE_TOP, CONTENT_W, RULE_Y - TITLE_TOP, anchor=MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    _set_para(p, line=1.06)
    run = p.add_run()
    run.text = title
    _set_font(run, size=_title_pt(title), bold=True, color=ctx.ink)
    # short accent rule under the title block
    _rect(slide, MARGIN_L, RULE_Y, 0.95, 0.045, ctx.accent)


def _footer(slide, ctx: RenderCtx, index: int, total: int, page=None) -> None:
    _rect(slide, MARGIN_L, FOOTER_Y, CONTENT_W, 0.012, _rgb(C_HAIRLINE))
    left = ctx.course
    ev = page.get("evidence") if isinstance(page, dict) else None
    if isinstance(ev, list) and ev:
        names = []
        for e in ev[:1]:
            if isinstance(e, dict) and e.get("file_name"):
                loc = str(e.get("page_or_slide") or "").strip()
                names.append(str(e["file_name"]) + (f" p{loc}" if loc else ""))
        if names:
            left = (left + "  ·  来源：" + "；".join(names)).strip(" ·")
    _simple_text(slide, left, MARGIN_L, FOOTER_Y + 0.04, CONTENT_W - 1.4, 0.32,
                 size=10, color=_rgb(C_MUTED))
    _simple_text(slide, f"{index:02d} / {total:02d}", SLIDE_W_IN - MARGIN_R - 1.4, FOOTER_Y + 0.04,
                 1.4, 0.32, size=10, color=_rgb(C_MUTED), align=PP_ALIGN.RIGHT)


def _bg(slide, color=C_BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


# --------------------------------------------------------------------------- #
# matplotlib helpers (formula / chart) + images
# --------------------------------------------------------------------------- #
def _latex_png(expr: str, tmpdir: str, *, color: str = C_INK, fontsize: int = 26) -> str | None:
    expr = str(expr or "").strip()
    if not expr:
        return None
    tex = expr if (expr.startswith("$") and expr.endswith("$")) else f"${expr}$"
    path = os.path.join(tmpdir, f"f_{uuid4().hex[:8]}.png")
    try:
        fig = plt.figure(figsize=(0.1, 0.1))
        fig.text(0.5, 0.5, tex, fontsize=fontsize, color=color, ha="center", va="center")
        fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.08, transparent=True)
        plt.close(fig)
        return path if os.path.exists(path) else None
    except Exception:
        plt.close("all")
        return None


def _chart_png(chart: dict, tmpdir: str, accent_hex: str) -> str | None:
    if not isinstance(chart, dict):
        return None
    ctype = str(chart.get("chartType") or chart.get("type") or "bar").lower()
    data = chart.get("data") if isinstance(chart.get("data"), dict) else {}
    labels = [str(x) for x in (data.get("labels") or [])]
    datasets = data.get("datasets") if isinstance(data.get("datasets"), list) else None
    path = os.path.join(tmpdir, f"c_{uuid4().hex[:8]}.png")
    try:
        fig, ax = plt.subplots(figsize=(7.6, 4.0))
        if datasets:
            for ds in datasets:
                vals = [float(v) for v in (ds.get("data") or []) if _is_num(v)]
                xs = labels[: len(vals)] or list(range(len(vals)))
                if ctype == "line":
                    ax.plot(xs, vals, marker="o", linewidth=2.2, color=accent_hex, label=str(ds.get("label", "")))
                else:
                    ax.bar(xs, vals, color=accent_hex, label=str(ds.get("label", "")))
            if any(ds.get("label") for ds in datasets):
                ax.legend(frameon=False)
        else:
            vals = [float(v) for v in (data.get("values") or []) if _is_num(v)]
            if not vals:
                plt.close(fig)
                return None
            xs = labels[: len(vals)] or list(range(len(vals)))
            if ctype == "pie":
                ax.pie(vals, labels=xs, autopct="%1.0f%%")
            elif ctype == "line":
                ax.plot(xs, vals, marker="o", linewidth=2.2, color=accent_hex)
            else:
                ax.bar(xs, vals, color=accent_hex)
        ax.tick_params(labelsize=11, colors="#5b6770")
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
        for spine in ("left", "bottom"):
            ax.spines[spine].set_color("#cdd6da")
        fig.tight_layout()
        fig.savefig(path, dpi=150, bbox_inches="tight", transparent=True)
        plt.close(fig)
        return path if os.path.exists(path) else None
    except Exception:
        plt.close("all")
        return None


def _is_num(v) -> bool:
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _resolve_image(page: dict, tmpdir: str) -> str | None:
    b64 = page.get("image_base64")
    if isinstance(b64, str) and b64.strip():
        raw = b64.split(",", 1)[1] if b64.strip().startswith("data:") else b64.strip()
        try:
            path = os.path.join(tmpdir, f"img_{uuid4().hex[:8]}.png")
            with open(path, "wb") as fh:
                fh.write(base64.b64decode(raw))
            return path
        except Exception:
            return None
    for key in ("image_path", "path", "image"):
        val = page.get(key)
        if isinstance(val, str) and os.path.exists(val):
            return val
    return None


def _fit(img_path: str, max_w: float, max_h: float) -> tuple[float, float]:
    try:
        w, h = Image.open(img_path).size
        if not w or not h:
            return max_w, max_h
        ratio = min(max_w / w, max_h / h)
        return w * ratio, h * ratio
    except Exception:
        return max_w, max_h


# --------------------------------------------------------------------------- #
# lightweight code syntax highlighting
# --------------------------------------------------------------------------- #
_KEYWORDS = {
    "python": {"def", "class", "return", "if", "elif", "else", "for", "while", "in", "import",
               "from", "as", "with", "try", "except", "finally", "raise", "and", "or", "not",
               "is", "None", "True", "False", "lambda", "yield", "pass", "break", "continue", "global"},
    "javascript": {"function", "return", "if", "else", "for", "while", "const", "let", "var",
                   "class", "new", "this", "import", "export", "from", "default", "try", "catch",
                   "throw", "async", "await", "true", "false", "null", "undefined", "typeof"},
    "java": {"public", "private", "protected", "class", "interface", "void", "int", "double",
             "float", "boolean", "char", "long", "return", "if", "else", "for", "while", "new",
             "static", "final", "this", "import", "package", "try", "catch", "throw", "true", "false", "null"},
    "cpp": {"int", "double", "float", "char", "bool", "void", "return", "if", "else", "for",
            "while", "class", "struct", "public", "private", "new", "delete", "const", "auto",
            "include", "using", "namespace", "std", "true", "false", "nullptr", "template"},
    "c": {"int", "double", "float", "char", "void", "return", "if", "else", "for", "while",
          "struct", "const", "static", "sizeof", "include", "define", "true", "false", "NULL"},
}
_KEYWORDS["js"] = _KEYWORDS["javascript"]
_KEYWORDS["c++"] = _KEYWORDS["cpp"]

_TOKEN_RE = re.compile(
    r"(?P<comment>#.*$|//.*$)|(?P<string>\"[^\"\n]*\"|'[^'\n]*')|"
    r"(?P<number>\b\d+\.?\d*\b)|(?P<word>[A-Za-z_]\w*)|(?P<ws>\s+)|(?P<other>.)"
)


def _hl_runs(line: str, lang: str) -> list[tuple[str, str]]:
    kws = _KEYWORDS.get((lang or "").lower(), set())
    runs: list[tuple[str, str]] = []
    for m in _TOKEN_RE.finditer(line):
        kind = m.lastgroup
        text = m.group()
        if kind == "comment":
            runs.append((text, C_CODE_COM))
        elif kind == "string":
            runs.append((text, C_CODE_STR))
        elif kind == "number":
            runs.append((text, C_CODE_NUM))
        elif kind == "word":
            runs.append((text, C_CODE_KW if text in kws else C_CODE_DEFAULT))
        else:
            runs.append((text, C_CODE_DEFAULT))
    return runs or [(line or " ", C_CODE_DEFAULT)]


# --------------------------------------------------------------------------- #
# content helpers
# --------------------------------------------------------------------------- #
def _bullet_block(slide, ctx, items, *, top, height, marker="square", base_pt=None):
    """Vertically-centred bullet list that fills the content zone."""
    items = [str(x).strip() for x in (items or []) if str(x).strip()]
    if not items:
        return
    n = len(items)
    pt = base_pt or (20 if n <= 4 else 18 if n == 5 else 17 if n <= 7 else 15)
    gap = 12 if n <= 4 else 9 if n <= 6 else 6
    box, tf = _textbox(slide, MARGIN_L, top, CONTENT_W, height, anchor=MSO_ANCHOR.MIDDLE)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_para(p, space_after=gap, line=1.18)
        if marker == "number":
            head = p.add_run()
            head.text = f"{i + 1:02d}   "
            _set_font(head, size=pt - 1, bold=True, color=ctx.accent, name=FONT_MONO)
        else:
            head = p.add_run()
            head.text = "▪  " if marker == "square" else "✓  "
            _set_font(head, size=pt, bold=True, color=ctx.accent)
        run = p.add_run()
        run.text = item
        _set_font(run, size=pt, color=_rgb(C_BODY))


# --------------------------------------------------------------------------- #
# per-type slide builders
# --------------------------------------------------------------------------- #
def _render_cover(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    # left accent spine + soft corner shape for identity
    _rect(slide, 0, 0, 0.22, SLIDE_H_IN, ctx.accent)
    _round_rect(slide, SLIDE_W_IN - 3.0, SLIDE_H_IN - 3.0, 4.2, 4.2, _rgb(C_SOFT))
    title = re.sub(r"\s+", " ", str(page.get("title") or "").strip())
    _simple_text(slide, "教学课件", MARGIN_L + 0.15, 2.35, 6, 0.4, size=13, bold=True, color=ctx.accent)
    box, tf = _textbox(slide, MARGIN_L + 0.15, 2.78, 9.6, 1.9, anchor=MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    _set_para(p, line=1.1)
    run = p.add_run()
    run.text = title
    _set_font(run, size=40 if len(title) <= 16 else 34 if len(title) <= 24 else 30, bold=True, color=ctx.ink)
    _rect(slide, MARGIN_L + 0.18, 4.75, 1.4, 0.06, ctx.accent)
    sub = str(page.get("subtitle") or "").strip()
    if sub:
        _simple_text(slide, sub, MARGIN_L + 0.15, 5.0, 9.6, 0.5, size=17, color=_rgb(C_MUTED))
    _simple_text(slide, "TeachMind · AI 智能备课", MARGIN_L + 0.15, SLIDE_H_IN - 0.7, 8, 0.4,
                 size=12, color=_rgb(C_MUTED))


def _render_agenda(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "目录", page.get("title", "课程导航"))
    items = [str(x) for x in (page.get("items") or page.get("points") or []) if str(x).strip()]
    zone_h = CONTENT_BOTTOM - CONTENT_TOP
    n = max(1, len(items))
    row_h = min(0.92, zone_h / n)
    block_h = row_h * n
    start = CONTENT_TOP + max(0, (zone_h - block_h) / 2)
    pt = 19 if n <= 6 else 16
    for i, item in enumerate(items[:9]):
        y = start + i * row_h
        _simple_text(slide, f"{i + 1:02d}", MARGIN_L, y, 0.8, row_h, size=pt + 2, bold=True,
                     color=ctx.accent, name=FONT_MONO, anchor=MSO_ANCHOR.MIDDLE)
        _simple_text(slide, item, MARGIN_L + 0.95, y, CONTENT_W - 0.95, row_h, size=pt,
                     color=_rgb(C_INK), anchor=MSO_ANCHOR.MIDDLE)
        if i < min(len(items), 9) - 1:
            _rect(slide, MARGIN_L, y + row_h, CONTENT_W, 0.01, _rgb(C_HAIRLINE))
    _footer(slide, ctx, idx, total, page)


def _render_content(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "讲解", page.get("title", ""))
    tip = str(page.get("tip") or "").strip()
    bottom = CONTENT_BOTTOM - (0.7 if tip else 0)
    bullets = page.get("bullets") or page.get("points") or page.get("takeaways") or []
    _bullet_block(slide, ctx, bullets, top=CONTENT_TOP, height=bottom - CONTENT_TOP, marker="square")
    if tip:
        _round_rect(slide, MARGIN_L, CONTENT_BOTTOM - 0.6, CONTENT_W, 0.56, _rgb(C_SOFT))
        _simple_text(slide, f"教学提示   {tip}", MARGIN_L + 0.25, CONTENT_BOTTOM - 0.55, CONTENT_W - 0.5,
                     0.46, size=12.5, italic=True, color=ctx.accent, anchor=MSO_ANCHOR.MIDDLE)
    _footer(slide, ctx, idx, total, page)


def _render_summary(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "小结", page.get("title", "课堂小结"))
    items = page.get("takeaways") or page.get("bullets") or []
    _bullet_block(slide, ctx, items, top=CONTENT_TOP, height=CONTENT_BOTTOM - CONTENT_TOP, marker="check")
    _footer(slide, ctx, idx, total, page)


def _render_code(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "代码", page.get("title", "代码示例"))
    lang = str(page.get("language") or "text")
    code = str(page.get("code") or "")
    expl = str(page.get("explanation") or "").strip()
    panel_bottom = CONTENT_BOTTOM - (0.5 if expl else 0)
    panel_h = panel_bottom - CONTENT_TOP
    _round_rect(slide, MARGIN_L, CONTENT_TOP, CONTENT_W, panel_h, _rgb(C_CODE_BG))
    _simple_text(slide, lang.upper(), SLIDE_W_IN - MARGIN_R - 1.5, CONTENT_TOP + 0.12, 1.35, 0.3,
                 size=10, bold=True, color=_rgb(C_ACCENT), align=PP_ALIGN.RIGHT, name=FONT_MONO)
    lines = code.split("\n")
    usable_h = panel_h - 0.55  # padding above/below code inside the panel
    n = max(1, len(lines))
    csize = 12.5 if n <= 12 else 10.5 if n <= 16 else 9.0 if n <= 20 else 8.0
    line_h_in = csize * 1.5 / 72.0  # conservative real line height incl. spacing
    max_lines = max(6, int(usable_h / line_h_in))
    if n > max_lines:
        lines = lines[: max_lines - 1] + ["# …（完整代码见随附文件）"]
    box, tf = _textbox(slide, MARGIN_L + 0.35, CONTENT_TOP + 0.3, CONTENT_W - 0.7, usable_h)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_para(p, line=1.3)
        for text, color in _hl_runs(line, lang):
            run = p.add_run()
            run.text = text
            _set_font(run, size=csize, color=_rgb(color), name=FONT_MONO)
    if expl:
        _simple_text(slide, expl, MARGIN_L, panel_bottom + 0.08, CONTENT_W, 0.42, size=13,
                     color=_rgb(C_BODY))
    _footer(slide, ctx, idx, total, page)


def _render_formula(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "公式", page.get("title", "公式推导"))
    formulas = page.get("formulas")
    if not isinstance(formulas, list) or not formulas:
        _bullet_block(slide, ctx, page.get("bullets") or [], top=CONTENT_TOP,
                      height=CONTENT_BOTTOM - CONTENT_TOP, marker="square")
        _footer(slide, ctx, idx, total, page)
        return
    formulas = formulas[:3]
    zone_h = CONTENT_BOTTOM - CONTENT_TOP
    slot = zone_h / len(formulas)
    for i, fm in enumerate(formulas):
        if isinstance(fm, dict):
            label = str(fm.get("label") or "")
            expr = str(fm.get("expr") or fm.get("latex") or "")
            fexpl = str(fm.get("explanation") or "")
        else:
            label, expr, fexpl = "", str(fm), ""
        y0 = CONTENT_TOP + i * slot
        _round_rect(slide, MARGIN_L, y0 + 0.06, CONTENT_W, slot - 0.18, _rgb(C_SOFT if i % 2 == 0 else C_SOFT2))
        if label:
            _simple_text(slide, label, MARGIN_L + 0.3, y0 + 0.16, CONTENT_W - 0.6, 0.34,
                         size=13, bold=True, color=ctx.accent)
        png = _latex_png(expr, ctx.tmpdir, fontsize=26)
        img_top = y0 + (0.5 if label else 0.2)
        img_max_h = slot - (0.95 if fexpl else 0.7)
        if png:
            w, h = _fit(png, CONTENT_W - 1.4, max(0.4, img_max_h))
            slide.shapes.add_picture(png, Inches(MARGIN_L + 0.4), Inches(img_top), width=Inches(w), height=Inches(h))
        else:
            _simple_text(slide, expr, MARGIN_L + 0.4, img_top, CONTENT_W - 0.8, 0.5, size=18,
                         color=_rgb(C_INK), name=FONT_MONO)
        if fexpl:
            _simple_text(slide, fexpl, MARGIN_L + 0.3, y0 + slot - 0.42, CONTENT_W - 0.6, 0.34,
                         size=11.5, italic=True, color=_rgb(C_MUTED))
    _footer(slide, ctx, idx, total, page)


def _render_two_column(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "对比", page.get("title", "对比"))
    left = page.get("left") if isinstance(page.get("left"), dict) else {}
    right = page.get("right") if isinstance(page.get("right"), dict) else {}
    gap = 0.5
    col_w = (CONTENT_W - gap) / 2
    zone_h = CONTENT_BOTTOM - CONTENT_TOP

    def column(x: float, data: dict, fill: str) -> None:
        head = str(data.get("title") or data.get("heading") or "")
        points = [str(p) for p in (data.get("points") or data.get("items") or []) if str(p).strip()]
        _round_rect(slide, x, CONTENT_TOP, col_w, zone_h, _rgb(fill))
        _rect(slide, x, CONTENT_TOP, col_w, 0.62, ctx.accent)
        _simple_text(slide, head, x + 0.25, CONTENT_TOP, col_w - 0.5, 0.62, size=16, bold=True,
                     color=_rgb("#FFFFFF"), anchor=MSO_ANCHOR.MIDDLE)
        box, tf = _textbox(slide, x + 0.3, CONTENT_TOP + 0.8, col_w - 0.6, zone_h - 1.0, anchor=MSO_ANCHOR.MIDDLE)
        pt = 16 if len(points) <= 5 else 13
        for i, point in enumerate(points[:8]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _set_para(p, space_after=8, line=1.18)
            head_run = p.add_run()
            head_run.text = "·  "
            _set_font(head_run, size=pt, bold=True, color=ctx.accent)
            run = p.add_run()
            run.text = point
            _set_font(run, size=pt, color=_rgb(C_BODY))

    column(MARGIN_L, left, C_SOFT)
    column(MARGIN_L + col_w + gap, right, C_SOFT2)
    _footer(slide, ctx, idx, total, page)


def _render_example(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "例题", page.get("title", "例题"))
    problem = str(page.get("problem") or "")
    steps = [str(s) for s in (page.get("steps") or []) if str(s).strip()]
    answer = str(page.get("answer") or "")
    y = CONTENT_TOP
    if problem:
        ph = 0.45 + 0.28 * (len(problem) // 38)
        _round_rect(slide, MARGIN_L, y, CONTENT_W, ph, _rgb(C_SOFT))
        _rect(slide, MARGIN_L, y, 0.08, ph, ctx.accent)
        _simple_text(slide, f"题目   {problem}", MARGIN_L + 0.3, y, CONTENT_W - 0.5, ph, size=15,
                     color=_rgb(C_INK), anchor=MSO_ANCHOR.MIDDLE)
        y += ph + 0.25
    if steps:
        sh = (CONTENT_BOTTOM - (0.55 if answer else 0)) - y
        box, tf = _textbox(slide, MARGIN_L + 0.1, y, CONTENT_W - 0.1, sh, anchor=MSO_ANCHOR.MIDDLE)
        for i, step in enumerate(steps[:7]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _set_para(p, space_after=8, line=1.2)
            head = p.add_run()
            head.text = f"{i + 1}   "
            _set_font(head, size=15, bold=True, color=ctx.accent, name=FONT_MONO)
            run = p.add_run()
            run.text = step
            _set_font(run, size=15, color=_rgb(C_BODY))
    if answer:
        _round_rect(slide, MARGIN_L, CONTENT_BOTTOM - 0.5, CONTENT_W, 0.48, _rgb(C_SOFT2))
        _simple_text(slide, f"答案   {answer}", MARGIN_L + 0.25, CONTENT_BOTTOM - 0.5, CONTENT_W - 0.5,
                     0.48, size=14, bold=True, color=ctx.accent, anchor=MSO_ANCHOR.MIDDLE)
    _footer(slide, ctx, idx, total, page)


def _render_chart(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "数据", page.get("title", "数据图表"))
    chart = page.get("chart") if isinstance(page.get("chart"), dict) else None
    if chart is None and (page.get("chartType") or page.get("data")):
        chart = {"chartType": page.get("chartType"), "data": page.get("data")}
    accent_hex = "#%02X%02X%02X" % (ctx.accent[0], ctx.accent[1], ctx.accent[2])
    png = _chart_png(chart, ctx.tmpdir, accent_hex) if chart else None
    if png:
        w, h = _fit(png, CONTENT_W - 1.0, CONTENT_BOTTOM - CONTENT_TOP - 0.6)
        slide.shapes.add_picture(png, Inches((SLIDE_W_IN - w) / 2), Inches(CONTENT_TOP + 0.1),
                                 width=Inches(w), height=Inches(h))
    else:
        _bullet_block(slide, ctx, page.get("bullets") or [], top=CONTENT_TOP,
                      height=CONTENT_BOTTOM - CONTENT_TOP, marker="square")
    cap = str(page.get("caption") or "")
    if cap:
        _simple_text(slide, cap, MARGIN_L, CONTENT_BOTTOM - 0.4, CONTENT_W, 0.36, size=12,
                     italic=True, color=_rgb(C_MUTED), align=PP_ALIGN.CENTER)
    _footer(slide, ctx, idx, total, page)


def _render_image(slide, page, ctx: RenderCtx, idx, total) -> None:
    _bg(slide)
    _header(slide, ctx, "图示", page.get("title", ""))
    path = _resolve_image(page, ctx.tmpdir)
    if path:
        w, h = _fit(path, CONTENT_W - 0.6, CONTENT_BOTTOM - CONTENT_TOP - 0.6)
        slide.shapes.add_picture(path, Inches((SLIDE_W_IN - w) / 2), Inches(CONTENT_TOP + 0.1),
                                 width=Inches(w), height=Inches(h))
    else:
        _bullet_block(slide, ctx, page.get("bullets") or ["（图片缺失）"], top=CONTENT_TOP,
                      height=CONTENT_BOTTOM - CONTENT_TOP, marker="square")
    cap = str(page.get("caption") or "")
    if cap:
        _simple_text(slide, cap, MARGIN_L, CONTENT_BOTTOM - 0.4, CONTENT_W, 0.36, size=12,
                     italic=True, color=_rgb(C_MUTED), align=PP_ALIGN.CENTER)
    _footer(slide, ctx, idx, total, page)


_BUILDERS = {
    "cover": _render_cover,
    "agenda": _render_agenda,
    "content": _render_content,
    "summary": _render_summary,
    "code": _render_code,
    "formula": _render_formula,
    "two_column": _render_two_column,
    "example": _render_example,
    "chart": _render_chart,
    "image": _render_image,
}


def _notes(slide, page) -> None:
    notes = str(page.get("notes") or "").strip()
    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = notes
        except Exception:
            pass


def render_pptx(intent: dict, slides_json: dict, output_path: str) -> str:
    pages = slides_json.get("pages") if isinstance(slides_json, dict) else None
    if not isinstance(pages, list) or not pages:
        raise ValueError("slides_json has no pages to render")

    theme = slides_json.get("theme") if isinstance(slides_json.get("theme"), dict) else {}
    ctx = RenderCtx(
        tmpdir=tempfile.mkdtemp(prefix="fastppt_render_"),
        accent=_rgb(C_ACCENT),  # use the curated palette rather than generator theme
        ink=_rgb(C_INK),
        course=str((intent or {}).get("topic") or "").strip(),
    )
    _ = theme
    try:
        prs = Presentation()
        prs.slide_width = SLIDE_W
        prs.slide_height = SLIDE_H
        total = len(pages)
        for i, page in enumerate(pages):
            if not isinstance(page, dict):
                continue
            page_type = str(page.get("type") or "content")
            builder = _BUILDERS.get(page_type, _render_content)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            try:
                builder(slide, page, ctx, i + 1, total)
            except Exception as exc:
                print(f"[pptx_renderer] page '{page.get('title')}' ({page_type}) failed: {exc}")
                try:
                    _render_content(slide, page, ctx, i + 1, total)
                except Exception:
                    pass
            _notes(slide, page)

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        prs.save(output_path)
        return output_path
    finally:
        shutil.rmtree(ctx.tmpdir, ignore_errors=True)
